from copy import deepcopy

import pytest

from ashare_value.demo import fixture_market
from ashare_value.domain import quality_metrics, rank_market
from ashare_value.pipeline import load_rules
from ashare_value.providers import parse_financial_abstract, parse_target_market


def enriched():
    market = fixture_market()
    for obs in market:
        obs.target_market = {"as_of": obs.as_of, "market_cap": 11e9, "dividend_yield": .04}
        for f in obs.financials:
            f.update(gross_margin=.30, net_margin=.15)
    return market


def test_six_equal_targets_and_strict_boundaries():
    obs = enriched()[0]
    rules = load_rules()
    assert quality_metrics(obs, rules)["target_count"] == 6
    obs.target_market.update(market_cap=10e9, dividend_yield=.03)
    obs.prices[-1]["pe"] = 20
    for f in obs.financials:
        f.update(gross_margin=.25, net_margin=.10, roe=.10)
    q = quality_metrics(obs, rules)
    assert q["passed"] and q["target_count"] == 1 and q["target_known_count"] == 6
    obs.prices[-1]["pe"] = 0
    assert quality_metrics(obs, rules)["target_count"] == 0
    obs.prices[-1]["pe"] = -1
    assert quality_metrics(obs, rules)["target_count"] == 0


def test_missing_is_unknown_not_a_smaller_denominator():
    obs = enriched()[0]
    obs.target_market = {"as_of": "2000-01-01", "market_cap": 11e9, "dividend_yield": .05}
    q = quality_metrics(obs, load_rules())
    assert q["passed"] and q["target_known_count"] == 4 and q["target_count"] == 4
    assert q["score"] == pytest.approx(400/6)
    assert sum(t["status"] == "待核查" for t in q["targets"]) == 2


def test_latest_complete_annual_not_half_year_or_future():
    obs = enriched()[0]
    obs.financials[-1].update(roe=.01, gross_margin=.01, net_margin=.01)
    obs.financials.append({**obs.financials[0], "period": "2026-12-31", "available_at": "2027-03-20", "roe": .01})
    q = quality_metrics(obs, load_rules())
    assert q["target_period"] == "2025-12-31" and q["target_count"] == 6


def test_soft_cash_debt_and_roe_keep_partial_matches():
    market = enriched()
    for obs in market:
        for f in obs.financials:
            f.update(roe=.05, ocf=1e6, debt_ratio=.8)
    q = quality_metrics(market[0], load_rules())
    assert q["passed"] and q["target_count"] == 5
    assert len(q["penalties"]) >= 2
    assert any(r["status"] == "候选" for r in rank_market(market, load_rules()))
    legacy = {**load_rules(), "quality_mode": "legacy"}
    assert not quality_metrics(market[0], legacy)["passed"]


def test_nonpositive_cumulative_profit_never_becomes_zero_conversion_peer():
    market = enriched()
    for f in market[0].financials:
        if f["period"] == "2021-12-31":
            f["net_profit"] = -1e10
    q = quality_metrics(market[0], load_rules())
    assert q["cash_conversion"] is None
    results = rank_market(market, load_rules())
    own = next(r for r in results if r["code"] == market[0].stock.code)
    assert own["status"] == "数据不足"
    assert all(market[0].stock.code not in r["peer"].get("codes", []) for r in results)


def test_target_count_dominates_secondary_score():
    market = enriched()
    # Company 0 is cheaper, but meets one fewer target than company 1.
    market[0].target_market["dividend_yield"] = .01
    for obs in market[2:]:
        obs.prices[-1]["adjusted_close"] = 100
    results = rank_market(market, load_rules())
    a, b = (next(r for r in results if r["code"] == market[i].stock.code) for i in (0, 1))
    assert a["status"] == b["status"] == "候选"
    assert b["quality"]["target_count"] == a["quality"]["target_count"] + 1
    assert b["score"] > a["score"] and b["rank"] < a["rank"]


def quote():
    return [{"SECURITY_CODE": "600062", "TRADE_DATE": "2026-09-09", "CLOSE_PRICE": 10, "TOTAL_MARKET_CAP": 12e9, "TOTAL_SHARES": 1.2e9}]


def dividend():
    return {"SECURITY_CODE": "600062", "EX_DIVIDEND_DATE": "2026-07-28", "REPORT_DATE": "2025-12-31",
            "ASSIGN_PROGRESS": "实施分配", "PRETAX_BONUS_RMB": 4, "NOTICE_DATE": "2026-07-22", "IMPL_PLAN_PROFILE": "10派4.00元(含税)"}


def test_dividend_uses_per_ten_shares_and_only_past_implemented():
    d = dividend()
    data = [d, {**d, "EX_DIVIDEND_DATE": "2026-10-01"}, {**d, "EX_DIVIDEND_DATE": "2025-09-09"}, {**d, "ASSIGN_PROGRESS": "董事会预案"}]
    value = parse_target_market(quote(), data, "sh.600062", "2026-09-09")
    assert value["market_cap"] == 12e9 and value["dividend_yield"] == pytest.approx(.04)
    assert len(value["dividend_items"]) == 1


@pytest.mark.parametrize("change", [{"IT_RATIO": 1}, {"IMPL_PLAN_PROFILE": "10派40元"}, {"NOTICE_DATE": "2026-10-01"}])
def test_ambiguous_dividend_cannot_score(change):
    value = parse_target_market(quote(), [{**dividend(), **change}], "sh.600062", "2026-09-09")
    assert value["market_cap"] == 12e9 and value["dividend_yield"] is None and value["errors"]


def test_duplicate_missing_and_no_dividend_distinct():
    duplicate = parse_target_market(quote(), [dividend(), deepcopy(dividend())], "sh.600062", "2026-09-09")
    missing = parse_target_market(quote(), None, "sh.600062", "2026-09-09")
    none = parse_target_market(quote(), [], "sh.600062", "2026-09-09")
    assert duplicate["dividend_yield"] is missing["dividend_yield"] is None
    assert none["dividend_yield"] == 0


def test_margins_use_percent_units_and_reject_conflicting_duplicates():
    rows = [{"指标": "毛利率", "20251231": 25.2}, {"指标": "销售净利率", "20251231": 10.3}]
    value = parse_financial_abstract(rows, ["20251231"], "2026-09-10")[0]
    assert value["gross_margin"] == pytest.approx(.252) and value["net_margin"] == pytest.approx(.103)
    with pytest.raises(ValueError):
        parse_financial_abstract(rows + [{"指标": "毛利率", "20251231": 26}], ["20251231"], "2026-09-10")


def test_market_cap_requires_same_date_and_total_share_reconciliation():
    for change in ({"TRADE_DATE": "2026-09-08"}, {"TOTAL_MARKET_CAP": 1.2e9}):
        value = parse_target_market([{**quote()[0], **change}], [dividend()], "sh.600062", "2026-09-09")
        assert value["market_cap"] is None and value["dividend_yield"] is None and value["errors"]


def test_targets_export_as_numbers_with_annual_date_and_explanation(tmp_path):
    from pathlib import Path

    from openpyxl import load_workbook

    from ashare_value.reports import export_reports
    from ashare_value.storage import Store

    market, rules = enriched(), load_rules()
    store = Store(tmp_path)
    try:
        results = rank_market(market, rules)
        identifier = store.save(market, results, rules, {"as_of": market[0].as_of}, kind="synthetic")
        output = export_reports(store, identifier)
        sheet = load_workbook(output["xlsx"])["研究候选"]
        cells = {sheet.cell(1, i).value: sheet.cell(2, i).value for i in range(1, sheet.max_column+1)}
        assert cells["六项目标达标数"] == 6 and cells["目标数据已知项数"] == 6
        assert cells["总市值（元）"] == 11e9 and cells["最近年度毛利率"] == .3
        assert cells["近12个月税前现金股息率"] == .04
        assert cells["目标财务年度"].strftime("%Y-%m-%d") == "2025-12-31"
        html = Path(output["html"]).read_text(encoding="utf-8")
        assert "六项质量目标：6/6 达标" in html and "达标数优先" in html and "2025-12-31" in html
    finally:
        store.close()
