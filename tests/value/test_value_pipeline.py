import sqlite3

import pytest
from openpyxl import load_workbook

from ashare_value.demo import fixture_market, save_demo
from ashare_value.domain import rank_market
from ashare_value.pipeline import DataClient, load_rules, market_scope
from ashare_value.providers import parse_financial_abstract
from ashare_value.reports import export_reports
from ashare_value.research import process_queue, safe_evidence_url, validate_adjustment
from ashare_value.storage import Store, process_lock


@pytest.fixture
def store(tmp_path):
    s = Store(tmp_path)
    yield s
    s.close()


def test_snapshot_observation_and_slot_are_immutable(store):
    identifier = save_demo(store, load_rules())
    for table in ("snapshots", "observations"):
        with pytest.raises(sqlite3.IntegrityError):
            store.db.execute(f"DELETE FROM {table}")
    with pytest.raises(sqlite3.IntegrityError):
        store.db.execute("UPDATE snapshots SET as_of='2020-01-01'")
    assert store.get(identifier)["kind"] == "synthetic"
    assert store.db.execute("SELECT count(*) FROM ai_tasks").fetchone()[0] == 0


def test_equal_windows_clock_timestamps_keep_latest_snapshot_and_previous_list(store, monkeypatch):
    monkeypatch.setattr("ashare_value.storage.utc_now", lambda: "2026-09-09T00:00:00+00:00")
    previous = save_demo(store, load_rules())
    current = save_demo(store, load_rules())
    assert store.get()["id"] == current
    assert store.changes(store.get())["previous"] == previous


def test_excel_and_html_from_same_snapshot_and_escape(store):
    market = fixture_market()
    market[0].stock.name = '<script>alert(1)</script>'
    rules = load_rules()
    identifier = store.save(market, rank_market(market, rules), rules, {"as_of":market[0].as_of}, kind="synthetic")
    out = export_reports(store, identifier)
    workbook = load_workbook(out["xlsx"])
    assert workbook["研究候选"].max_row > 1
    assert workbook["研究候选"]["F2"].data_type == "n"
    assert workbook["研究候选"]["H2"].number_format == "0.0%"
    from pathlib import Path
    text = Path(out["html"]).read_text(encoding="utf-8")
    assert '<script>alert(1)</script>' not in text
    assert '&lt;script&gt;' in text
    again = export_reports(store, identifier)
    assert again["html"] != out["html"]


def test_offline_client_does_not_access_network(store):
    with pytest.raises(RuntimeError, match="离线缓存缺失"):
        DataClient(store, load_rules(), offline=True).request("universe")


def test_exchange_scope_excludes_shenzhen_and_beijing_before_peer_selection():
    from dataclasses import replace
    stock = fixture_market()[0].stock
    stocks = [replace(stock, code=code) for code in ("sh.600600", "sh.688001", "sz.000001", "bj.920001")]
    assert [s.code for s in market_scope(stocks, {"exchanges": ["sh"]})] == ["sh.600600", "sh.688001"]
    assert len(market_scope(stocks, {"exchanges": ["sh", "sz"]})) == 3


def test_shanghai_scan_scope_and_report_disclosure(store, monkeypatch):
    from dataclasses import asdict, replace
    from datetime import datetime
    from zoneinfo import ZoneInfo

    from ashare_value.pipeline import scan
    stock = fixture_market()[0].stock
    records = [asdict(replace(stock, code=code, industry_code="J66")) for code in ("sh.600000", "sz.000001", "bj.920001")]
    today = datetime.now(ZoneInfo("Asia/Shanghai")).date().isoformat()
    responses = {"universe": {"stocks": records, "industry_date": today, "limitations": []}, "market_date": today, "calendar": []}
    responses["peer_groups"] = {"groups": {"sh.600000": {"name": "验证银行", "peer_group": "银行"}}, "total": 1}
    monkeypatch.setattr(DataClient, "request", lambda self, operation, **kwargs: responses[operation])
    identifier = scan(store, load_rules(), progress=lambda *a, **k: None)
    snapshot = store.get(identifier)
    assert [r["code"] for r in snapshot["results"]] == ["sh.600000"]
    assert snapshot["coverage"]["source_universe_count"] == 3
    assert snapshot["coverage"]["universe_count"] == 1
    assert snapshot["coverage"]["outside_scope_count"] == 2
    from pathlib import Path
    report = export_reports(store, identifier)
    assert "目标：上交所 A 股" in Path(report["html"]).read_text(encoding="utf-8")


def test_parallel_offline_collection_uses_separate_connections_and_actual_date(store):
    from datetime import datetime
    from zoneinfo import ZoneInfo

    from ashare_value.pipeline import collect_observations
    from ashare_value.storage import digest
    market = fixture_market()[:4]
    today = datetime.now(ZoneInfo("Asia/Shanghai")).date().isoformat()
    for obs in market:
        for operation, arguments, payload in (
            ("prices", {"code": obs.stock.code, "end": obs.as_of}, {"prices": obs.prices}),
            ("financials", {"code": obs.stock.code, "available_at": today}, {"financials": obs.financials}),
        ):
            store.cache(digest({"adapter": "1.0.0", "operation": operation, "arguments": arguments}), payload)
        obs.prices, obs.financials, obs.research_date = [], [], "2000-01-01"
    rules = load_rules()
    rules["request_pause_seconds"] = 0
    collect_observations(store, rules, market, True, lambda *a, **k: None)
    assert all(obs.prices and obs.financials and not obs.errors and obs.research_date == today for obs in market)


def test_midnight_financial_request_retries_with_new_date(store, monkeypatch):
    from ashare_value.pipeline import fetch_observation
    dates = iter(["2026-09-09", "2026-09-10"])
    monkeypatch.setattr("ashare_value.pipeline.research_day", lambda: next(dates))
    seen = []
    def request(self, operation, **kwargs):
        if operation == "prices":
            return {"prices": []}
        seen.append(kwargs["available_at"])
        if kwargs["available_at"] == "2026-09-09":
            raise RuntimeError("请求跨午夜完成")
        return {"financials": [{"available_at": "2026-09-10"}]}
    monkeypatch.setattr(DataClient, "request", request)
    obs = fixture_market()[0]
    fetch_observation(store.root, load_rules(), obs)
    assert seen == ["2026-09-09", "2026-09-10"]
    assert obs.research_date == "2026-09-10" and not obs.errors
    assert obs.financials[0]["available_at"] == "2026-09-10"


def test_financial_ratio_units_and_conflicting_duplicates():
    fields = [{"指标":"资产负债率", "20251231": 61.17}, {"指标":"净资产收益率(ROE)", "20251231": 15.0}]
    row = parse_financial_abstract(fields, ["20251231"], "2026-09-09")[0]
    assert row["debt_ratio"] == .6117 and row["roe"] == .15
    fields.append({"指标":"资产负债率", "20251231": .6117})
    with pytest.raises(ValueError, match="重复字段"):
        parse_financial_abstract(fields, ["20251231"], "2026-09-09")


@pytest.mark.parametrize("duplicate", [False, True])
def test_business_group_pagination_rejects_partial_or_duplicate_lists(monkeypatch, duplicate):
    from ashare_value.providers import get_peer_groups
    class Response:
        def __init__(self, page):
            self.page = page
        def raise_for_status(self):
            pass
        def json(self):
            offset = 0 if duplicate else (self.page-1)*100
            return {"data": {"total": 2001, "diff": [{"f12": str(600000+i), "f13": 1, "f14": "测试公司", "f100": "测试业务"} for i in range(offset, min(offset+100, 2001))]}}
    monkeypatch.setattr("requests.get", lambda url, params, timeout: Response(params["pn"]))
    if duplicate:
        with pytest.raises(ValueError, match="重复"):
            get_peer_groups()
    else:
        result = get_peer_groups()
        assert result["total"] == len(result["groups"]) == 2001


def test_ai_failure_preserves_snapshot_and_retries(store, monkeypatch):
    rules = load_rules()
    market = fixture_market()
    identifier = store.save(market, rank_market(market, rules), rules, {"as_of":market[0].as_of})
    def fail(*args, **kwargs):
        raise RuntimeError("额度不足")
    monkeypatch.setattr("ashare_value.research.invoke_codex", fail)
    result = process_queue(store, rules, progress=lambda *a, **k: None)
    assert result["completed"] == 0 and result["pending"] > 0
    assert store.db.execute("SELECT count(*) FROM ai_tasks WHERE status='retry'").fetchone()[0] == 1
    assert store.get(identifier)["results"] == rank_market(market, rules)


def test_adjustments_without_evidence_cannot_be_applied():
    result = validate_adjustment({"period":"2025-12-31", "after_tax_parent_impact":-1e8, "source_id":"fake"}, {}, {"2025-12-31"})
    assert result["status"] == "pending" and result["numeric_validation"] == "未通过"


@pytest.mark.parametrize("url", ["https://cninfo.com.cn.evil.test/a", "http://127.0.0.1/a", "file:///C:/secret", "https://user@cninfo.com.cn/a"])
def test_research_cannot_fetch_arbitrary_private_or_fake_sources(url):
    assert not safe_evidence_url(url)


def test_os_lock_prevents_overlapping_runs(tmp_path):
    with process_lock(tmp_path):
        with pytest.raises(RuntimeError, match="已有扫描"):
            with process_lock(tmp_path):
                pytest.fail("overlap")


def test_completed_schedule_slot_does_not_rescan(tmp_path, monkeypatch):
    from ashare_value.cli import main
    from ashare_value.pipeline import latest_due
    s = Store(tmp_path)
    identifier = save_demo(s, load_rules())
    with s.db:
        s.db.execute("INSERT INTO schedule_slots VALUES(?,?)", (latest_due(), identifier))
    s.close()
    def unexpected(*args, **kwargs):
        pytest.fail("completed slot attempted network scan")
    monkeypatch.setattr("ashare_value.cli.scan", unexpected)
    assert main(["--data-dir", str(tmp_path), "due", "--no-ai"]) == 0


def test_rule_change_keeps_old_slot_and_runs_new_version(tmp_path, monkeypatch):
    from ashare_value.cli import main
    from ashare_value.pipeline import latest_due
    s = Store(tmp_path)
    rules = load_rules()
    rules["version"] = "旧规则测试"
    previous = save_demo(s, rules)
    with s.db:
        s.db.execute("INSERT INTO schedule_slots VALUES(?,?)", (latest_due(), previous))
    s.close()
    monkeypatch.setattr("ashare_value.cli.scan", lambda store, rules, **kwargs: save_demo(store, rules))
    assert main(["--data-dir", str(tmp_path), "due", "--no-ai"]) == 0
    s = Store(tmp_path)
    try:
        assert s.db.execute("SELECT count(*) FROM schedule_slots").fetchone()[0] == 2
        assert s.db.execute("SELECT snapshot_id FROM schedule_slots WHERE slot=?", (latest_due(),)).fetchone()[0] == previous
    finally:
        s.close()


def test_source_failure_still_outputs_reports_without_completing_slot(tmp_path, monkeypatch):
    from ashare_value.cli import main
    def failure(*args, **kwargs):
        raise RuntimeError("模拟免费数据源不可用")
    monkeypatch.setattr("ashare_value.pipeline.DataClient.request", failure)
    assert main(["--data-dir", str(tmp_path), "due", "--no-ai"]) == 2
    s = Store(tmp_path)
    try:
        assert s.db.execute("SELECT count(*) FROM schedule_slots").fetchone()[0] == 0
        assert s.get()["coverage"]["errors"]
        assert list(tmp_path.glob("reports/*/*/*.xlsx")) and list(tmp_path.glob("reports/*/*/*.html"))
    finally:
        s.close()
