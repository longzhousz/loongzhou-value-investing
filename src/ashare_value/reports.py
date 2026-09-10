"""独立 Python 程序的本地 Excel / HTML 导出；所有输出来自正式快照。"""
from __future__ import annotations

import html
import json
import math
import textwrap
from datetime import date, datetime, timezone
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CHECKS = "已检查：非ST状态、历史完整性、报表盈利、长期ROE、现金流支撑、负债率、收入盈利趋势、同窗价格和估值低位、质量可比同行。未全面排除价值陷阱；审计、关联交易、担保诉讼、资产回收等仍需材料核查。"
METHOD = "价格位置=(当期后复权收盘−窗口最低)/(窗口最高−窗口最低)，不是价格分位；PE/PB 使用每日历史报表口径估值的中秩经验分位。负PE/PB和缺失不算低估；当前价格为对应交易日未复权实际收盘。3年、5年分别评估，同一窗口价格与PE/PB同时通过且同行不贵才入选。"


def display(value):
    if value is None:
        return "不可用"
    if isinstance(value, list) and all(isinstance(v, str) for v in value):
        return "；".join(value) if value else "无"
    if isinstance(value, list) and value and all(isinstance(v, dict) and "reason" in v and "points" in v for v in value):
        return "；".join(f"{v['reason']}（扣 {v['points']} 分）" for v in value)
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, indent=2)
    return str(value)


def escape(value):
    return html.escape(display(value), quote=True)


COLUMNS = [
    ("综合排序", "rank"), ("代码", "code"), ("名称", "name"), ("行业", "industry"), ("状态", "status"),
    ("实际收盘价（元）", "price"), ("价格交易日", "as_of"), ("3年价格位置", "windows.3.price_position"),
    ("5年价格位置", "windows.5.price_position"), ("PE(TTM)", "pe"), ("PB(MRQ)", "pb"),
    ("3年PE分位", "windows.3.pe_percentile"), ("3年PB分位", "windows.3.pb_percentile"),
    ("5年PE分位", "windows.5.pe_percentile"), ("5年PB分位", "windows.5.pb_percentile"),
    ("同行PE位置", "peer.pe_percentile"), ("同行PB位置", "peer.pb_percentile"), ("可比同行数", "peer.count"),
    ("长期ROE中位数", "quality.median_roe"), ("累计现金转化率", "quality.cash_conversion"),
    ("最新资产负债率", "quality.debt_ratio"), ("基本面质量分", "quality.score"), ("最新盈利同比", "quality.profit_yoy"),
    ("财务期间", "quality.latest_period"), ("实际价格覆盖年限", "coverage_years"),
    ("数据可信度", "confidence"), ("综合分", "score"), ("入选或排除原因", "reasons"),
    ("扣分原因", "penalties"), ("风险", "risks"), ("标签", "tags"), ("剩余未知", "unknown"),
    ("正常化情景", "normalized"), ("事件影响", "events"), ("AI状态", "ai_status"), ("来源及采集日期", "sources"), ("同行业务分组", "peer_group"),
    ("六项目标达标数", "quality.target_count"), ("目标数据已知项数", "quality.target_known_count"),
    ("总市值（元）", "quality.market_cap"), ("最近年度毛利率", "quality.gross_margin"),
    ("最近年度净利率", "quality.net_margin"), ("近12个月税前现金股息率", "quality.dividend_yield"),
    ("最近完整年度ROE", "quality.annual_roe"), ("目标财务年度", "quality.target_period"),
    ("六项质量目标明细", "quality.targets"), ("同达标档次级分", "secondary_score")]

AI_STATUS = {"pending": "待研究", "retry": "等待重试", "needs_review": "研究完成，仍有待核查项", "reviewed": "已复核", "superseded": "旧规则待办已停止"}


def target_value(metric, value):
    if value is None:
        return "不可用"
    if metric in ("gross_margin", "net_margin", "dividend_yield", "annual_roe"):
        return f"{value:.3%}"
    return f"{value:,.2f}"


def nested(row, path):
    value = row
    for key in path.split("."):
        if not isinstance(value, dict):
            return None
        value = value.get(key)
    return value


def excel_value(value):
    if isinstance(value, (dict, list)):
        value = display(value)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        value = "'" + value
    if isinstance(value, str) and len(value) == 10 and value[4:5] == "-" and value[7:8] == "-":
        try:
            return date.fromisoformat(value)
        except ValueError:
            pass
    return value if value is not None else "不可用"


def table_sheet(book, name, headers, rows):
    sheet = book.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append([excel_value(v) for v in row])
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 34
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="163D35")
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F0F5F2")
            if isinstance(cell.value, float):
                cell.number_format = "0.00"
            elif isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"
        sheet.row_dimensions[row[0].row].height = 60
    for i, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(i)].width = 18 if len(header) < 12 else 25
    sheet.sheet_view.showGridLines = False
    return sheet


def fit_rows(sheet):
    for row in sheet.iter_rows(min_row=2):
        lines = 1
        for cell in row:
            width = sheet.column_dimensions[cell.column_letter].width or 18
            text = str(cell.value or "")
            wrapped = sum(max(1, math.ceil(sum(2 if ord(c) > 127 else 1 for c in line)/max(5, width-2))) for line in text.splitlines())
            lines = max(lines, wrapped)
        sheet.row_dimensions[row[0].row].height = min(409, max(24, lines*16+8))


def window_description(windows):
    parts = []
    for years, window in windows.items():
        parts.append(escape(f"{years}年：价格位置{pct(window.get('price_position'))}，PE分位{pct(window.get('pe_percentile'))}，PB分位{pct(window.get('pb_percentile'))}。{window['reason']}"))
    return "<br>".join(parts) or "暂无可用窗口"


def export_reports(store, identifier=None):
    snapshot = store.get(identifier)
    results = json.loads(json.dumps(snapshot["results"]))
    analyses = {}
    for row in store.db.execute("SELECT t.code,t.status,a.payload,a.created_at FROM ai_tasks t LEFT JOIN ai_analyses a ON a.task_id=t.id WHERE t.snapshot_id=? ORDER BY a.created_at,a.rowid", (snapshot["id"],)):
        analyses[row["code"]] = {"task_status": row["status"], "analysis": json.loads(row["payload"]) if row["payload"] else None}
    for item in results:
        entry = analyses.get(item["code"])
        item["ai_status"] = AI_STATUS.get(entry["task_status"], entry["task_status"]) if entry else "未排队"
        if entry and entry["analysis"]:
            analysis = entry["analysis"]
            item["normalized"] = analysis.get("normalized", [])
            item["events"] = analysis.get("events", [])
            item["unknown"].extend(analysis.get("pending", []))
            item["risks"].extend(analysis.get("risks", []))
            if analysis.get("event_label_validated") and item["status"] == "候选":
                item["tags"].append("事件驱动型价值低位")
    candidates = [r for r in results if r["status"] == "候选"]
    changes = store.changes(snapshot)
    folder = store.root / "reports" / snapshot["id"] / (datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ-") + uuid4().hex[:6])
    folder.mkdir(parents=True, exist_ok=False)
    coverage = snapshot["coverage"]
    scoring_method = ("六项等权软目标，质量目标分=达标数/6×100；综合分=15×达标数+0.099×同档次级分。次级分为风险质量50%、便宜35%、数据把握15%，仅三年历史再扣5分后截断为0～100；所以达标数优先。缺失不算达标。"
                      if snapshot["rules"].get("quality_mode") == "six_targets" else "旧模式：质量50%、便宜35%、数据把握15%，仅三年历史另扣5分。")
    title = "A股价值研究候选"
    if snapshot["kind"] != "live":
        title += "（离线合成验证，非真实候选）"
    book = Workbook()
    book.remove(book.active)
    summary_rows = [["报告类型", title], ["目标范围", coverage.get("target")], ["行情交易日", snapshot["as_of"]], ["采集与研究日期", coverage.get("research_date")],
                    ["来源原始名录公司数", coverage.get("source_universe_count", coverage.get("universe_count"))],
                    ["主源范围内公司数", coverage.get("primary_scope_count")], ["辅助行业名录数", coverage.get("auxiliary_universe_count")],
                    ["仅辅助名录有记录数", coverage.get("auxiliary_only_count")], ["同行细分行业来源", coverage.get("peer_group_source")],
                    ["范围外公司数", coverage.get("outside_scope_count")],
                    ["候选数量", len(candidates)], ["扫描方式", coverage.get("mode")], ["名录公司数", coverage.get("universe_count")],
                    ["本轮选定公司数", coverage.get("requested_count")], ["已记录公司数", coverage.get("observed_count")],
                    ["实际取得价格公司数", coverage.get("price_count")], ["实际取得财务公司数", coverage.get("financial_count")],
                    ["完整有效三年公司数", coverage.get("complete_3y_count")], ["完整有效五年公司数", coverage.get("complete_5y_count")],
                    ["覆盖状态", coverage.get("status_counts")], ["采集错误公司数", coverage.get("collection_failed_count")],
                    ["数据源限制", coverage.get("limitations")], ["运行错误", coverage.get("errors")], ["检查范围和未知", CHECKS],
                    ["计算口径", METHOD], ["质量目标与排序", scoring_method], ["规则版本", snapshot["rules"]["version"]], ["规则内容SHA256", snapshot["rule_hash"]],
                    ["名单变化", changes], ["全部规则", snapshot["rules"]]]
    summary = table_sheet(book, "覆盖与口径", ["项目", "内容"], summary_rows)
    summary.column_dimensions["A"].width = 29
    summary.column_dimensions["B"].width = 110
    summary.freeze_panes = "B2"
    detail_rows = []
    detail_links = {}
    for record in results:
        for label, key in COLUMNS:
            value = nested(record, key)
            if not isinstance(value, (dict, list, str)) or len(display(value)) <= 100:
                continue
            first_row = len(detail_rows) + 2
            detail_links[(record["code"], key)] = first_row
            lines = [part for line in display(value).splitlines() for part in (textwrap.wrap(line, width=40, replace_whitespace=False) or [""])]
            for offset in range(0, len(lines), 8):
                detail_rows.append([record["code"], record["name"], label if offset == 0 else label + "（续）", "\n".join(lines[offset:offset+8])])
    for name, records in [("研究候选", candidates), ("全部筛选记录", results)]:
        sheet = table_sheet(book, name, [h for h, _ in COLUMNS], [[nested(r, k) for _, k in COLUMNS] for r in records])
        for row_index, record in enumerate(records, 2):
            for col_index, (_, key) in enumerate(COLUMNS, 1):
                detail_row = detail_links.get((record["code"], key))
                if detail_row:
                    cell = sheet.cell(row_index, col_index)
                    preview = display(nested(record, key)).replace("\n", " ")[:54]
                    cell.value = excel_value(preview + "…\n点击查看完整明细")
                    cell.hyperlink = f"#'指标与证据明细'!D{detail_row}"
                    cell.font = Font(name="微软雅黑", size=10, color="175E49", underline="single")
        for index, (label, key) in enumerate(COLUMNS, 1):
            if any(term in label for term in ("分位", "价格位置", "同行PE位置", "同行PB位置", "ROE", "转化率", "负债率", "同比", "毛利率", "净利率", "股息率")):
                for column in sheet.iter_cols(min_col=index, max_col=index, min_row=2):
                    for cell in column:
                        cell.number_format = "0.0%"
                        if key in ("quality.gross_margin", "quality.net_margin", "quality.dividend_yield", "quality.annual_roe"):
                            cell.number_format = "0.000%"
            if key in ("reasons", "penalties", "risks", "unknown", "sources", "normalized", "events"):
                sheet.column_dimensions[get_column_letter(index)].width = 55
    evidence_rows = []
    for code, entry in analyses.items():
        evidence_rows.append([code, "任务状态", AI_STATUS.get(entry["task_status"], entry["task_status"])])
        analysis = entry["analysis"] or {}
        evidence_rows.append([code, "研究摘要", analysis.get("summary", "尚未完成材料研究")])
        for key, label in [("risks", "风险"), ("pending", "待核查"), ("sources", "来源与原文核验"), ("events", "事件影响"), ("adjustments", "盈利调整建议"), ("normalized", "正常化情景")]:
            evidence_rows.extend([code, label, value] for value in analysis.get(key, []))
    ai_sheet = table_sheet(book, "AI证据与核查", ["代码", "研究项目", "内容及依据"], evidence_rows)
    ai_sheet.column_dimensions["C"].width = 110
    details = table_sheet(book, "指标与证据明细", ["代码", "名称", "字段", "完整内容（长文本按续行保存）"], detail_rows)
    details.column_dimensions["D"].width = 100
    xlsx = folder / "候选与筛选记录.xlsx"
    for sheet in book:
        fit_rows(sheet)
    book.save(xlsx)
    sections = []
    for item in candidates:
        targets = item["quality"].get("targets", [])
        target_rows = "".join(f"<tr><td>{escape(t['label'])}</td><td>{target_value(t['metric'], t['value'])}</td><td>{escape(t['operator'])} {target_value(t['metric'], t['threshold'])}</td><td>{escape(t['status'])}</td><td>{escape(t['period'])}</td></tr>" for t in targets)
        target_panel = f"<h3>六项质量目标：{item['quality'].get('target_count')}/6 达标</h3><p>等权软目标，未达标仍可保留；缺失不算达标。先按达标数排序，同档再比较风险、便宜程度和数据把握。毛利率、净利率及ROE使用最近完整年度。达标按原始精度判断，显示值仅四舍五入。</p><table><tr><th>指标</th><th>实际值</th><th>目标</th><th>状态</th><th>日期/期间</th></tr>{target_rows}</table>" if targets else ""
        windows = "".join(f"<tr><td>{y} 年</td><td>{pct(w.get('price_position'))}</td><td>{pct(w.get('pe_percentile'))}</td><td>{pct(w.get('pb_percentile'))}</td><td>{escape(w['reason'])}</td></tr>" for y, w in item["windows"].items())
        ai = analyses.get(item["code"], {})
        sections.append(f"""<article><h2>{item['rank']}. {escape(item['name'])} <small>{escape(item['code'])}</small></h2>
        <p>{escape(item['industry'])} · 实际收盘 {escape(item['price'])} 元（{escape(item['as_of'])}） · 综合分 {item['score']} · 数据可信度：{item['confidence']}</p>
        <h3>公司是否仍然好</h3><p>长期ROE {pct(item['quality'].get('median_roe'))}；累计现金转化率 {pct(item['quality'].get('cash_conversion'))}；资产负债率 {pct(item['quality'].get('debt_ratio'))}。</p>
        {target_panel}
        <p>扣分与风险：{escape(item['risks'])}</p><h3>当前价格是否足够低</h3><p>PE(TTM) {item['pe']:.2f}，PB(MRQ) {item['pb']:.2f}。同行分组：{escape(item.get('peer_group'))}。同行PE位置 {pct(item['peer'].get('pe_percentile'))}，PB位置 {pct(item['peer'].get('pb_percentile'))}，可比同行 {item['peer'].get('count')} 家。</p>
        <table><tr><th>窗口</th><th>价格位置</th><th>PE分位</th><th>PB分位</th><th>判断</th></tr>{windows}</table>
        <h3>判断有多大把握</h3><p>入选理由：{escape(item['reasons'])}</p><p>标签：{escape(item['tags'])}</p><p>待核查：{escape(item['unknown'])}</p>
        <details><summary>正常化盈利与事件分析（不改变基础资格和排序）</summary><pre>{escape(item['normalized'])}\n{escape(item['events'])}</pre></details>
        <details><summary>AI意见、证据和来源日期</summary><pre>{escape(ai)}</pre></details>
        <details><summary>客观数据来源</summary><pre>{escape(item['sources'])}</pre></details></article>""")
    audit = "".join(f"<tr><td>{escape(r['code'])}</td><td>{escape(r['name'])}</td><td>{escape(r['industry'])}</td><td>{escape(r['status'])}</td><td>{escape(r['reasons'])}</td><td>{window_description(r['windows'])}</td></tr>" for r in results)
    supplemental = "".join(f"<article><h2>补充材料研究：{escape(code)}</h2><p>状态：{escape(AI_STATUS.get(entry['task_status'], entry['task_status']))}。补充研究不改变基础资格和排序。</p><pre>{escape(entry['analysis'])}</pre></article>" for code, entry in analyses.items())
    document = f"""<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
    <meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'; img-src data:; base-uri 'none'">
    <title>{escape(title)}</title><style>body{{font:16px/1.75 'Microsoft YaHei',sans-serif;color:#173b33;background:#f3f5ef;margin:0}}main{{max-width:1180px;margin:auto;padding:38px 28px}}h1{{font-size:32px}}h2{{font-size:24px}}small{{font-size:14px;color:#65746c}}article,.panel{{background:white;padding:28px;margin:22px 0;border:1px solid #d5dfd7;border-radius:8px}}.count{{font-size:40px}}table{{border-collapse:collapse;width:100%;font-size:14px}}td,th{{padding:10px;border-bottom:1px solid #d5dfd7;text-align:left;vertical-align:top}}th{{background:#e5eee7}}pre{{white-space:pre-wrap;overflow-wrap:anywhere;font:14px/1.6 'Microsoft YaHei',sans-serif}}details{{margin:14px 0}}.scroll{{overflow:auto}}a{{color:#175e49}}@media print{{body{{background:white}}article{{break-inside:avoid}}}}</style>
    <main><p>本地价值研究 / {escape(snapshot['as_of'])}</p><h1>{escape(title)}</h1>
    <p>规则决定资格，证据决定把握，标签解释原因，排序决定先研究谁。</p>
    <div class="panel"><span class="count">{len(candidates)}</span> 家候选。{'符合条件不足 5 家，如实保留，不凑数。' if len(candidates)<5 else '按基础规则优先级排列，最多 20 家。'}
    <p>目标：{escape(coverage.get('target', '快照未记录'))}；范围内名录 {coverage.get('universe_count',0)} 家，选定 {coverage.get('requested_count',0)} 家，取得财务 {coverage.get('financial_count',0)} 家，完整有效三年 {coverage.get('complete_3y_count',0)} 家、五年 {coverage.get('complete_5y_count',0)} 家。</p>
    <p>来源原始名录 {coverage.get('source_universe_count', coverage.get('universe_count',0))} 家；范围外 {escape(coverage.get('outside_scope_count'))} 家。候选及同行范围以本快照规则为准。</p>
    <p>{escape(coverage.get('limitations',[]))}</p><p>覆盖状态：{escape(coverage.get('status_counts',{}))}</p><p>运行错误：{escape(coverage.get('errors',[]))}</p>
    <p><a href="候选与筛选记录.xlsx">打开 Excel 候选和完整筛选表</a></p></div>
    {''.join(sections) if candidates else '<article><h2>本次没有正式入选公司</h2><p>请结合下方排除记录和数据完整性判断原因。没有候选不代表市场没有机会。</p></article>'}
    {supplemental}
    <article><h2>名单变化</h2><pre>{escape(changes)}</pre><h2>方法和剩余未知</h2><p>{METHOD}</p><p>{escape(scoring_method)}</p><p>{CHECKS}</p><p>正常化盈利保留报表盈利和实际资产负债影响；尚无逐时点/同行可比正常化序列时，不用正常化PE代替历史分位，不给予事件奖励。</p>
    <details><summary>全部可修改规则与版本</summary><pre>{escape(snapshot['rules'])}</pre><p>规则SHA256：{snapshot['rule_hash']}</p></details>
    <details><summary>覆盖范围和数据缺失详情</summary><pre>{escape(coverage)}</pre></details></article>
    <article><h2>全部筛选记录</h2><div class="scroll"><table><tr><th>代码</th><th>名称</th><th>行业</th><th>状态</th><th>原因</th><th>窗口详情</th></tr>{audit}</table></div></article>
    <p>快照 {escape(snapshot['id'])}，生成于 {escape(datetime.now().astimezone().isoformat())}。用途为研究排序，不预测短期涨跌，不连接交易。</p></main></html>"""
    report = folder / "研究报告.html"
    report.write_text(document, encoding="utf-8")
    (folder / "report.json").write_text(json.dumps({"snapshot_id": snapshot["id"], "coverage": coverage, "rules": snapshot["rules"], "results": results, "changes": changes, "ai": analyses}, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    return {"snapshot_id": snapshot["id"], "html": str(report), "xlsx": str(xlsx), "candidate_count": len(candidates)}


def pct(value):
    return f"{value:.1%}" if isinstance(value, (float, int)) else "不可用"
